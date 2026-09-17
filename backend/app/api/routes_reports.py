import io
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from ..services.analysis import get_analysis
from ..services.report_generator import build_pdf

try:
    import openpyxl
    _OPENPYXL_OK = True
except ImportError:
    _OPENPYXL_OK = False

router = APIRouter()

@router.get("/report/{analysis_id}", response_class=StreamingResponse)
def report_pdf(analysis_id: str):
    b = get_analysis(analysis_id)
    if not b:
        raise HTTPException(status_code=404, detail="Unknown analysis_id")
    pdf = build_pdf(b)
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename=StatStock_{b['symbol']}_{analysis_id}.pdf"})

@router.get("/export/{analysis_id}")
def export_csv(analysis_id: str, kind: str = "raw"):
    import pandas as pd
    from fastapi.responses import StreamingResponse
    b = get_analysis(analysis_id)
    if not b:
        raise HTTPException(status_code=404, detail="Unknown analysis_id")
    if kind.lower() in ("xlsx", "excel"):
        return export_excel(analysis_id)
    df = pd.DataFrame(b["records"])
    buf = io.StringIO(); df.to_csv(buf, index=False)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv",
                             headers={"Content-Disposition": f"attachment; filename={b['symbol']}_{kind}.csv"})


@router.get("/export/{analysis_id}/excel", response_class=StreamingResponse)
def export_excel(analysis_id: str):
    if not _OPENPYXL_OK:
        raise HTTPException(status_code=500, detail="Excel export needs the 'openpyxl' package. Restart the backend after running: pip install -r requirements.txt")
    from fastapi.responses import StreamingResponse
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(status_code=500, detail="Excel export needs the 'openpyxl' package. Restart the backend after running: pip install -r requirements.txt")

    b = get_analysis(analysis_id)
    if not b:
        raise HTTPException(status_code=404, detail="Unknown analysis_id")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="111113", end_color="111113", fill_type="solid")
    title_font = Font(bold=True, size=13)
    meta_font = Font(size=10, color="555555")
    ccy = (b.get("currency") or "INR").upper()
    headers = ["Date", f"Open ({ccy})", f"High ({ccy})", f"Low ({ccy})", f"Close ({ccy})", "Volume"]
    widths = [14, 14, 14, 14, 14, 18]

    def write_raw_sheet(ws, title: str, subtitle: str, records: list[dict]):
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws["A1"] = title
        ws["A1"].font = title_font
        ws["A2"] = subtitle
        ws["A2"].font = meta_font
        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=4, column=col, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")
        for i, r in enumerate(records, start=5):
            ws.cell(row=i, column=1, value=r["date"]).number_format = "yyyy-mm-dd"
            for col, key in ((2, "open"), (3, "high"), (4, "low"), (5, "close")):
                ws.cell(row=i, column=col, value=float(r[key])).number_format = '#,##0.00'
            ws.cell(row=i, column=6, value=int(r["volume"])).number_format = '#,##0'
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A5"
        if records:
            ws.auto_filter.ref = f"A4:F{4 + len(records)}"


    nifty_records = b.get("nifty_records")
    if not nifty_records and b.get("nifty_dates"):
        nifty_records = [{"date": d, "open": c, "high": c, "low": c, "close": c, "volume": 0}
                         for d, c in zip(b["nifty_dates"], b["nifty_close"])]

    wb = Workbook()
    ws_stock = wb.active
    ws_stock.title = "Stock Raw Data"
    write_raw_sheet(ws_stock,
                    f"{b['stock_name']} ({b['symbol']}) — raw daily prices, {b['start']} to {b['end']}",
                    f"Data source: {b['provider']} · Retrieved: {b['retrieved_at']} · {b['n']} trading days",
                    b["records"])
    ws_nifty = wb.create_sheet("Nifty Raw Data")
    write_raw_sheet(ws_nifty,
                    f"NIFTY 50 — raw daily prices, {b['start']} to {b['end']}",
                    f"Data source: {b['provider']} · Retrieved: {b['retrieved_at']} · {b['n_nifty']} trading days",
                    nifty_records or [])

    info = wb.create_sheet("Info")
    info["A1"] = "Field"
    info["B1"] = "Value"
    for c in ("A1", "B1"):
        info[c].font = header_font
        info[c].fill = header_fill
    rows = [("Stock", f"{b['stock_name']} ({b['symbol']})"), ("Exchange", b["exchange"]),
            ("Currency", ccy),
            ("Period", f"{b['start']} to {b['end']}"), ("Price field", b.get("price_field", "Close")),
            ("Data source", b["provider"]), ("Retrieved at", b["retrieved_at"]),
            ("Trading days", b["n"]), ("Last close", f"{b['last_close']} on {b['last_date']}")]
    for i, (k, v) in enumerate(rows, start=2):
        info.cell(row=i, column=1, value=k).font = Font(bold=True)
        info.cell(row=i, column=2, value=v)
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 70

    buf = io.BytesIO()
    try:
        wb.save(buf)
    except Exception:
        raise HTTPException(status_code=500, detail="Could not build the Excel file. Please try again in a moment.")
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=StatStock_{b['symbol']}_{analysis_id}.xlsx"})
